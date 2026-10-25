"""Validate document package, navigation, source/evidence references and local extraction."""
from pathlib import Path
from datetime import datetime,timezone
from docx import Document
from lxml import etree
import json,re,hashlib,zipfile
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'docs/deliverables'
path=OUT/'HEDAMI-Pricing-Advisor-Technical-Guide.docx'
checks=[]
def check(name,condition):
    checks.append({'check':name,'passed':bool(condition)})
    assert condition,name
with zipfile.ZipFile(path) as z:
    check('DOCX ZIP integrity',z.testzip() is None)
    for name in z.namelist():
        if name.endswith(('.xml','.rels')):etree.fromstring(z.read(name))
    check('All XML and relationship parts parse',True)
    xml=etree.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
    anchors=xml.xpath('//w:bookmarkStart/@w:name',namespaces=ns)
    refs=xml.xpath('//w:hyperlink/@w:anchor',namespaces=ns)
    check('22 chapter bookmarks and matching internal navigation',len(anchors)==22 and set(refs)<=set(anchors))
    check('Six embedded images with descriptive alternative text',len(xml.xpath('//wp:docPr[@descr]',namespaces=ns))==6 and len([n for n in z.namelist() if n.startswith('word/media/')])==6)
    alltext=' '.join(xml.xpath('//w:t/text()',namespaces=ns))
    check('No unresolved content directives',not re.search(r'\{\{(?:SNIP|FILE|MANIFEST|FIG)',alltext))
    check('No API key-like credential pattern',not re.search(r'sk-[A-Za-z0-9_-]{12,}',alltext))
    check('All requested central concepts present',all(x in alltext for x in ['harness','RuntimeInstructions','PlanReview','agentic','phase','second phase','decimal','request_clarification']))
    check('No placeholder TODO markers',not re.search(r'\bTODO\b|TBD|Lorem ipsum',alltext))
    check('No false default 2013 creation timestamp',b'2013-12-23' not in z.read('docProps/core.xml') and b'2013-12-24' not in z.read('docProps/core.xml'))
    footers=''.join(z.read(n).decode() for n in z.namelist() if re.fullmatch(r'word/footer\d+\.xml',n))
    check('Page-number field present',' PAGE ' in footers)
source=(OUT/'technical-guide-final.md').read_text()
headings=re.findall(r'^# (.+)',source,re.M)
check('22 source chapters; second phase is final',len(headings)==22 and headings[-1].startswith('22. What a second phase'))
snapshot=json.loads((OUT/'source-snapshot.json').read_text())
check('All source snapshot hashes match current files',all(hashlib.sha256((ROOT/f['path']).read_bytes()).hexdigest()==f['sha256'] for f in snapshot['files']))
check('17 actual code excerpts within valid source line ranges',len(snapshot['snippets'])==17 and all(0<s['startLine']<=s['endLine']<=len((ROOT/s['file']).read_text().splitlines()) for s in snapshot['snippets']))
check('Required preserved reports and feedback trace exist',all((ROOT/p).exists() for p in ['docs/evidence/baseline/evaluations-20261024T043445508.json','docs/evidence/followup/evaluations-20261024T043908058.json','docs/evidence/followup/runs/85d6a4cc56a74a7b9051f0a378393c3a.json']))
case_counts={p.name:len(re.findall(r'\[Fact\]|\[InlineData\(',p.read_text())) for p in (ROOT/'tests/Hedami.PricingAdvisor.Domain.Tests').glob('*Tests.cs')}
check('Documented local case count matches inspected test declarations',sum(case_counts.values())==57)
d=Document(path);check('Word package opens through python-docx',len(d.inline_shapes)==6 and len(d.tables)==22)
extracted=Path('/tmp/hedami-technical-guide.txt')
check('macOS textutil extracted readable text',extracted.exists() and all(s in extracted.read_text() for s in ['What this solution is','What a second phase could do','0.44420565','InstructionsVersion']))
report={'validatedAtUtc':datetime.now(timezone.utc).isoformat(),'file':path.name,'fileSha256':hashlib.sha256(path.read_bytes()).hexdigest(),'checks':checks,'counts':{'sourceWords':len(source.split()),'chapters':len(headings),'figures':len(d.inline_shapes),'tablesIncludingCover':len(d.tables),'actualCodeExcerpts':len(snapshot['snippets']),'localCasesInspected':sum(case_counts.values())},'visualReview':{'technicalFigures':'Six-figure montage inspected; arrow/label improvements applied.','fullWordPagination':'Not visually rendered: macOS Quick Look failed sandbox initialization. No claim of full-page rendering. Word content/structure and text extraction were checked.'},'runtimeImpact':'Documentation/scripts/assets only; no application credential read and no paid model calls.'}
(OUT/'document-validation.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps({'passed':len(checks),'failed':0,'counts':report['counts'],'sha256':report['fileSha256']},indent=2))
