"""Build the draft and final Word guides from reviewed local Markdown/source assets."""
from pathlib import Path
from datetime import datetime, timezone
import hashlib,json,re,textwrap,zipfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/deliverables';ASSETS=OUT/'assets'
NAVY='17324D';TEAL='087F8C';LIGHT='E9F2F8';GRAY='F3F6F8'
SNIPS=[];SOURCES=set()
def expand(source):
    def snippet(m):
        path,needle,count=m.group(1).split('|');lines=(ROOT/path).read_text().splitlines()
        start=next(i for i,l in enumerate(lines) if needle in l);chosen=lines[start:start+int(count)]
        end=start+len(chosen)
        SNIPS.append({'file':path,'startLine':start+1,'endLine':end});SOURCES.add(path)
        code='\n'.join(f'{i+start+1:>4} | {line}' for i,line in enumerate(chosen))
        return f'Source excerpt: {path} — lines {start+1}–{end}.\n\n```csharp\n{code}\n```'
    source=re.sub(r'\{\{SNIP:([^}]+)\}\}',snippet,source)
    def procedure(m):
        path=m.group(1);SOURCES.add(path)
        return 'Procedure file: '+path+'\n\n'+ '\n\n'.join((ROOT/path).read_text().split('\n\n')[1:])
    source=re.sub(r'\{\{FILE:([^}]+)\}\}',procedure,source)
    for path in ['global.json','src/Hedami.PricingAdvisor.Console/Program.cs','src/Hedami.PricingAdvisor.Console/LiveEvaluations.cs','src/Hedami.PricingAdvisor.Console/LocalConfiguration.cs','src/Hedami.PricingAdvisor.Application/Harness/RunContracts.cs','src/Hedami.PricingAdvisor.Application/Serialization.cs','tests/Hedami.PricingAdvisor.Domain.Tests/Hedami.PricingAdvisor.Domain.Tests.csproj']:
        SOURCES.add(path)
    manifest=[{'path':p,'sha256':hashlib.sha256((ROOT/p).read_bytes()).hexdigest()} for p in sorted(SOURCES)]
    (OUT/'source-snapshot.json').write_text(json.dumps({'evidenceAsOf':'2026-10-24','files':manifest,'snippets':SNIPS},indent=2)+'\n')
    table='| Source file | SHA-256 prefix |\n| --- | --- |\n'+'\n'.join('| '+x['path']+' | '+x['sha256'][:16]+' |' for x in manifest)
    source=source.replace('{{MANIFEST}}',table+'\n\nThe complete hashes and snippet line ranges are preserved in docs/deliverables/source-snapshot.json. This snapshot does not assert a Git commit exists; the session could not create one.')
    return source

def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr();e=OxmlElement('w:shd');e.set(qn('w:fill'),fill);pr.append(e)
def hyper(par,text,target,anchor=False):
    el=OxmlElement('w:hyperlink')
    if anchor:el.set(qn('w:anchor'),target)
    else:el.set(qn('r:id'),par.part.relate_to(target,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),TEAL);pr.append(color)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
    r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);el.append(r);par._p.append(el)
def inline(par,text):
    pattern=r'(\*\*[^*]+\*\*|`[^`]+`|https?://[^\s;]+)'
    for piece in re.split(pattern,text):
        if not piece:continue
        if piece.startswith('**') and piece.endswith('**'):par.add_run(piece[2:-2]).bold=True
        elif piece.startswith('`') and piece.endswith('`'):
            r=par.add_run(piece[1:-1]);r.font.name='Courier New';r.font.size=Pt(9)
        elif piece.startswith(('http://','https://')):
            target=piece.rstrip('.,)');hyper(par,target,target)
            if len(target)!=len(piece):par.add_run(piece[len(target):])
        else:par.add_run(piece)
def field(par,name):
    r=par.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');r._r.append(a)
    r=par.add_run();instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' '+name+' ';r._r.append(instr)
    r=par.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'end');r._r.append(a)
def base(title,subtitle):
    d=Document();sec=d.sections[0]
    sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
    sec.top_margin=Inches(.70);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
    sec.header_distance=Inches(.28);sec.footer_distance=Inches(.28)
    sec.different_first_page_header_footer=True
    normal=d.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(7)
    normal.paragraph_format.widow_control=True
    for name,size in [('Title',32),('Subtitle',15),('Heading 1',22),('Heading 2',14),('Heading 3',12)]:
        st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(NAVY)
        st.paragraph_format.space_before=Pt(15);st.paragraph_format.space_after=Pt(8)
        st.paragraph_format.keep_with_next=True
    d.styles['Heading 1'].paragraph_format.space_before=Pt(0)
    if 'Code' not in d.styles:
        st=d.styles.add_style('Code',1);st.font.name='Courier New';st.font.size=Pt(8)
        st.paragraph_format.line_spacing=1.0;st.paragraph_format.space_after=Pt(1)
        st.paragraph_format.left_indent=Inches(.10);st.paragraph_format.right_indent=Inches(.08)
    st=d.styles['Caption'];st.font.name='Arial';st.font.size=Pt(9);st.font.color.rgb=RGBColor.from_string('506779')
    st.paragraph_format.space_after=Pt(11)
    hdr=sec.header.paragraphs[0];hdr.text='HEDAMI  |  Pricing Advisor';hdr.style='Caption';hdr.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer=sec.footer.paragraphs[0];footer.style='Caption';footer.add_run('Technical study guide  •  Evidence snapshot: 24 Oct 2026     |     ');field(footer,'PAGE');footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    settings=d.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'en-US');normal.element.get_or_add_rPr().append(lang)
    d.core_properties.title=title;d.core_properties.subject='Architecture, agentic loops, harness, runtime skills, code, evaluations and phase two'
    d.core_properties.author='Heron Dantas — prepared with Codex assistance'
    d.core_properties.last_modified_by='Codex documentation generator'
    timestamp=datetime.now(timezone.utc)
    d.core_properties.created=timestamp
    d.core_properties.modified=timestamp
    d.core_properties.revision=2 if 'Initial' not in subtitle else 1
    d.core_properties.keywords='HEDAMI; .NET; agentic loop; harness; tool use; runtime skills; pricing; evaluations'
    d.core_properties.comments='Generated from local source and inspected synthetic evidence. No private credentials included.'
    d.add_paragraph('HEDAMI 3D LABS',style='Subtitle')
    d.add_paragraph(title,style='Title');d.add_paragraph(subtitle,style='Subtitle')
    d.add_paragraph('Prepared for Heron Dantas',style='Heading 2')
    d.add_paragraph('Detailed study document and English portfolio reference\nEvidence snapshot: 24 October 2026\n'+('Version 0.1 — initial outline before self-review' if 'Initial' in subtitle else 'Version 1.0 — revised after documented author self-review'))
    box=d.add_table(rows=1,cols=1);box.alignment=WD_TABLE_ALIGNMENT.CENTER;shade(box.cell(0,0),LIGHT)
    box.cell(0,0).text='Purpose: understand and defend the implementation.\nStatus: executed console prototype; fictional business inputs; recommendation only.\nNo claim of production readiness or demonstrated superiority of extra agents.'
    d.add_paragraph('The code, evidence and limitations are explained together. Runtime procedures and development skills are distinguished explicitly. This document does not contain an API key.',style='Caption')
    return d

def add_table(d,rows):
    cols=len(rows[0]);t=d.add_table(rows=1,cols=cols);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    total=6.8277
    if cols==2:fractions=[.40,.60]
    elif cols==3:fractions=[.22,.35,.43]
    elif cols==4:fractions=[.21,.20,.25,.34]
    elif cols>=6:fractions=[.17]+[(.83)/(cols-1)]*(cols-1)
    else:fractions=[1/cols]*cols
    for i,col in enumerate(t.columns):col.width=Inches(total*fractions[i])
    allcells=[]
    for rid,row in enumerate(rows):
        cells=t.rows[0].cells if rid==0 else t.add_row().cells
        for i,text in enumerate(row):
            cell=cells[i];cell.width=Inches(total*fractions[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.04
            inline(p,text)
            for r in p.runs:r.font.name='Arial';r.font.size=Pt(8.7 if cols>=6 else 9)
            if rid==0:
                shade(cell,NAVY)
                for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
            elif rid%2:shade(cell,GRAY)
            margin=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');margin.append(e)
            cell._tc.get_or_add_tcPr().append(margin)
        trpr=t.rows[rid]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if rid==0:
            repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trpr.append(repeat)
    d.add_paragraph().paragraph_format.space_after=Pt(1)

def render(d,source,final=True):
    headings=[re.sub(r'^# ','',l) for l in source.splitlines() if l.startswith('# ')]
    d.add_page_break();d.add_paragraph('Chapter navigation',style='Heading 1')
    d.add_paragraph('Click a title below or use Word’s Navigation Pane. Page numbers are generated by Word during layout.')
    for i,h in enumerate(headings):
        p=d.add_paragraph();p.paragraph_format.space_after=Pt(3);hyper(p,h,'chapter_'+str(i),True)
    d.add_page_break();lines=source.splitlines();i=0;chapter=0;bookmark=1
    while i<len(lines):
        line=lines[i]
        if not line.strip():i+=1;continue
        if line.startswith('```'):
            language=line[3:];i+=1;code=[]
            while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
            for n,l in enumerate(code):
                wrapped=textwrap.wrap(l,width=96,expand_tabs=False,replace_whitespace=False,drop_whitespace=False,break_long_words=True,break_on_hyphens=False) or ['']
                for j,part in enumerate(wrapped):
                    p=d.add_paragraph(style='Code');p.add_run(part if j==0 else '       '+part)
                    p.paragraph_format.keep_with_next=(n<len(code)-1 or j<len(wrapped)-1)
                    pr=p._p.get_or_add_pPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),GRAY);pr.append(sh)
            d.add_paragraph().paragraph_format.space_after=Pt(3);i+=1;continue
        fig=re.fullmatch(r'\{\{FIG:([^|]+)\|(.+)\}\}',line)
        if fig:
            p=d.add_paragraph();p.paragraph_format.keep_with_next=True
            run=p.add_run();pic=run.add_picture(str(ASSETS/fig.group(1)),width=Inches(6.80))
            pic._inline.docPr.set('descr',fig.group(2));pic._inline.docPr.set('title',fig.group(1))
            d.add_paragraph(fig.group(2),style='Caption');i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                cells=[s.strip() for s in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c) for c in cells):rows.append(cells)
                i+=1
            assert len(set(map(len,rows)))==1,rows
            add_table(d,rows);continue
        m=re.match(r'^(#{1,3}) (.+)',line)
        if m:
            level=len(m.group(1));p=d.add_paragraph(m.group(2),style='Heading '+str(level))
            if level==1:
                if final and chapter>0:p.paragraph_format.page_break_before=True
                b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(bookmark));b.set(qn('w:name'),'chapter_'+str(chapter));p._p.insert(0,b)
                e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(bookmark));p._p.append(e);bookmark+=1;chapter+=1
            i+=1;continue
        if line.startswith('- '):p=d.add_paragraph(style='List Bullet');inline(p,line[2:]);i+=1;continue
        chunks=[line];i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','```','{{FIG:','- ')):
            chunks.append(lines[i]);i+=1
        text=' '.join(chunks);p=d.add_paragraph()
        if text.startswith(('Source excerpt:','Procedure file:')):
            p.style='Caption';p.paragraph_format.keep_with_next=True
        inline(p,text)
    return d

source=expand((OUT/'technical-guide-source.md').read_text())
assert '{{SNIP:' not in source and '{{FILE:' not in source and '{{MANIFEST}}' not in source
(OUT/'technical-guide-final.md').write_text(source)
final=base('HEDAMI Pricing Advisor','Solution, Architecture and Agentic Engineering Guide')
render(final,source,True)
final_path=OUT/'HEDAMI-Pricing-Advisor-Technical-Guide.docx';final.save(final_path)
draft=base('HEDAMI Pricing Advisor','Initial documentation outline — v0.1')
render(draft,(OUT/'documentation-outline-v0.1.md').read_text(),False)
draft_path=OUT/'HEDAMI-Pricing-Advisor-Initial-Outline.docx';draft.save(draft_path)
for path in [final_path,draft_path]:
    with zipfile.ZipFile(path) as z:assert z.testzip() is None
print(json.dumps({'final':str(final_path),'bytes':final_path.stat().st_size,'bodyWords':len(source.split()),'chapters':len(re.findall(r'^# ',source,re.M)),'figures':len(final.inline_shapes),'tables':len(final.tables),'sourceExcerpts':len(SNIPS),'draft':str(draft_path)},indent=2))
