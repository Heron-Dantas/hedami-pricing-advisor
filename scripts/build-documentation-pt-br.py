"""Generate the private Brazilian Portuguese study edition without changing English artifacts."""
from pathlib import Path
import ast,hashlib,json,re,zipfile
from docx import Document
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/deliverables/pt-BR';OUT.mkdir(parents=True,exist_ok=True)
english=ROOT/'docs/deliverables/HEDAMI-Pricing-Advisor-Technical-Guide.docx'
english_hash=hashlib.sha256(english.read_bytes()).hexdigest()
translations={
'docs/deliverables/assets':'docs/deliverables/pt-BR/assets',
'Architecture: local controls, external model':'Arquitetura: controles locais, modelo externo',
'Implemented console prototype | arrows show application calls and returned evidence':'Protótipo implementado no console | setas mostram chamadas e resultados',
'Human':'Pessoa','Objective / feedback':'Objetivo / feedback','Compose and render':'Montar e exibir',
'State, checks, loop and budgets':'Estado, validações, loop e limites',
'OpenAI model':'Modelo OpenAI','Analyst or fresh reviewer context':'Contexto do analista ou contexto novo do revisor',
'Runtime skills':'Skills','Two trusted files':'Dois arquivos confiáveis',
'Typed tools':'Ferramentas tipadas','Read, calculate, draft':'Consultar, calcular, rascunhar',
'C# domain':'Domínio C#','decimal arithmetic':'Aritmética com decimal',
'Local catalog':'Catálogo local','Synthetic JSON':'JSON fictício','Local traces':'Registros locais','Observable JSON':'JSON de eventos observáveis',
'request':'pedido','result':'resultado','calculate':'calcular','read':'consultar','events':'eventos',
'Main loop: model chooses, application executes':'Loop principal: modelo escolhe, aplicação executa',
'PricingHarness.RunAsync | success and failure paths have explicit stop conditions':'PricingHarness.RunAsync | sucesso e falha têm condições explícitas de parada',
'Objective':'Objetivo','Human request':'Pedido humano','Harness checks':'Verificações do harness',
'Context, usage and cancellation':'Contexto, consumo e cancelamento','Model response':'Resposta do modelo',
'Exactly one function call':'Exatamente uma chamada de função','Tool request':'Pedido de ferramenta','Name + JSON arguments':'Nome + argumentos JSON',
'Validate':'Validar','Schema, permission and plan':'Schema, permissão e plano','Execute tool':'Executar ferramenta','Stored data / calculator':'Dados armazenados / calculador',
'Record and append':'Registrar e adicionar','Result uses same call_id':'Resultado usa o mesmo call_id','terminal':'fim',
'continue with tool output':'continuar com o resultado da ferramenta',
'Draft / clarification / bounded stop':'Rascunho / esclarecimento / interrupção por limite',
'No arbitrary price application':'Nenhuma aplicação arbitrária de preço',
'Reflection and collaboration: separated contexts':'Reflexão e colaboração: contextos separados',
'One provider/model may serve both roles | no independent truth oracle':'Um modelo pode atender os dois papéis | revisor não é uma fonte infalível',
'Analyst context':'Contexto do analista','Objective, tool results and draft':'Objetivo, resultados de ferramentas e rascunho',
'Verify evidence and coordinate':'Verificar evidências e coordenar','Reviewer context':'Contexto do revisor',
'Objective + draft; no analyst history':'Objetivo + rascunho; sem histórico do analista',
'fresh context':'contexto novo','revise':'revisar','bounded feedback':'feedback limitado','approve':'aprovar',
'Verdict + traceable findings':'Parecer + achados rastreáveis','Revision budget':'Limite de revisões','One revised draft by default':'Um rascunho revisado por padrão',
'Reviewed draft or stop':'Rascunho revisado ou interrupção','Human approval still required':'Aprovação humana continua necessária',
'Adaptive planning: the selected action gates execution':'Planejamento adaptativo: a ação escolhida controla a execução',
'The planner is an analyst mode, not a third independent agent':'Planejamento é um modo do analista, não um terceiro agente independente',
'Analyst selects':'Analista seleciona','update_plan(nextTool, objective)':'update_plan(nextTool, objective)',
'Dependency checks':'Verificar dependências','Read before calculate; evidence before draft':'Consultar antes de calcular; evidência antes do rascunho',
'Stored NextTool':'NextTool armazenada','Short observable objective':'Objetivo curto e observável',
'Execution gate':'Controle da execução','Requested tool must match; selection is consumed':'Ferramenta deve corresponder; seleção é consumida',
'Tool result':'Resultado da ferramenta','Advance evidence or revise next action':'Avançar evidências ou rever a próxima ação',
'replan after results / errors':'replanejar após resultado / erro',
'Safe exit: request_clarification':'Saída segura: request_clarification',
'Missing information can stop without an artificial tool sequence':'Informação ausente permite encerrar sem uma sequência artificial',
'Human feedback: a new linked run with current data':'Feedback humano: novo fluxo vinculado com dados atuais',
'Demonstrated real workflow | prior target replaced by exactly 25%':'Fluxo real demonstrado | meta anterior substituída por exatamente 25%',
'Prior run':'Execução anterior','Draft or questions; immutable run ID':'Rascunho ou perguntas; ID imutável',
'Human feedback':'Feedback humano','New objective overrides conflicting prior targets':'Novo objetivo substitui metas anteriores conflitantes',
'Context projection':'Resumo do contexto','Compact prior evidence; parentRunId':'Evidência anterior resumida; parentRunId',
'New harness':'Novo harness','Fresh run / reread catalog and rules':'Nova execução / consultar catálogo e regras novamente',
'Recalculate':'Recalcular','Stored inputs; target 25%; discount 10%':'Dados armazenados; meta 25%; desconto 10%',
'New draft':'Novo rascunho','BRL 41.03 / 36.93; margins 25.01% / 17.79%':'R$ 41,03 / 36,93; margens 25,01% / 17,79%',
'No catalog mutation and no price application are introduced by feedback.':'Feedback não altera o catálogo nem aplica preços.',
'Phase two: prioritize evidence before expansion':'Segunda fase: priorizar evidência antes da expansão',
'PROPOSED ONLY | each gate needs measurable acceptance and human decisions':'APENAS PROPOSTA | cada etapa exige critérios e decisões humanas',
'Gate 1: quality and review':'Etapa 1: qualidade e revisão',
'Outcome-based grading; actionable critiques; retain historical failures':'Avaliar resultados; críticas executáveis; preservar falhas históricas',
'Gate 2: business inputs':'Etapa 2: dados de negócio',
'Per-field provenance; verified costs; actual settlement assumptions':'Procedência por campo; custos verificados; liquidação real',
'Gate 3: context and state':'Etapa 3: contexto e estado',
'Compact evidence; durable reservations; tested recovery and concurrency':'Evidência resumida; reservas duráveis; recuperação e concorrência testadas',
'Gate 4: human workflow':'Etapa 4: decisão humana',
'Usable review UI; governed edits; identity-bound approval if needed':'Interface útil; edição governada; aprovação ligada à identidade se necessária',
'Conditional later: Azure / MCP / authorized channel writes':'Depois, se justificado: Azure / MCP / escrita autorizada em canais'
}
class Localize(ast.NodeTransformer):
    def visit_Constant(self,node):
        if isinstance(node.value,str) and node.value in translations:return ast.copy_location(ast.Constant(translations[node.value]),node)
        return node
figure_source=(ROOT/'scripts/build-documentation-figures.py').read_text().replace('font(36,True);bf=font(30)','font(32,True);bf=font(28)').replace('font(48,True)','font(42,True)')
# Fit translated labels to the same geometry rather than clipping longer Portuguese text.
node_start=figure_source.index('    def node(')
node_end=figure_source.index('    def edge(',node_start)
node_code="""    def node(self,box,title,body='',color=BLUE):
        x1,y1,x2,y2=box
        self.d.rounded_rectangle(box,20,fill=color,outline=NAVY,width=3)
        def lines(text,face,maxw):
            out=[]
            for part in text.split('\\n'):
                line=''
                for word in part.split():
                    probe=(line+' '+word).strip()
                    if self.d.textlength(probe,font=face)>maxw and line:
                        out.append(line);line=word
                    else:line=probe
                out.append(line)
            return out
        for size in range(32,21,-1):
            f=font(size,True);bf=font(size-4)
            tl=lines(title,f,x2-x1-32);bl=lines(body,bf,x2-x1-32) if body else []
            th=size+8;bh=size+4;gap=10 if bl else 0
            height=len(tl)*th+len(bl)*bh+gap
            if height<=y2-y1-16:break
        assert height<=y2-y1-10,(title,box,height)
        y=y1+(y2-y1-height)/2
        for line in tl:
            self.d.text(((x1+x2-self.d.textlength(line,font=f))/2,y),line,font=f,fill=INK);y+=th
        y+=gap
        for line in bl:
            self.d.text(((x1+x2-self.d.textlength(line,font=bf))/2,y),line,font=bf,fill=INK);y+=bh
"""
figure_source=figure_source[:node_start]+node_code+figure_source[node_end:]
figure_ast=ast.fix_missing_locations(Localize().visit(ast.parse(figure_source)))
exec(compile(figure_ast,'localized-figures','exec'),{'__file__':str(ROOT/'scripts/build-documentation-figures.py')})
# Import only definitions: the original script's generation entrypoint is deliberately not executed.
writer_source=(ROOT/'scripts/build-word-documentation.py').read_text().split('\nsource=expand(',1)[0]
ns={'__file__':str(ROOT/'scripts/build-word-documentation.py')}
exec(compile(writer_source,'word-writer-definitions','exec'),ns)
ns['OUT']=OUT;ns['ASSETS']=OUT/'assets';ns['SNIPS']=[]
ns['SOURCES']={'runtime-skills/analyze-profitability.md','runtime-skills/review-pricing.md'}
source=ns['expand']((OUT/'guia-tecnico-fonte.md').read_text())
source=source.replace('Source excerpt:','Fonte do trecho:').replace(' — lines ',' — linhas ')
source=source.replace('| Source file | SHA-256 prefix |','| Arquivo de origem | Prefixo SHA-256 |')
source=source.replace('The complete hashes and snippet line ranges are preserved in docs/deliverables/source-snapshot.json. This snapshot does not assert a Git commit exists; the session could not create one.',
'Os hashes completos e intervalos dos trechos estão em docs/deliverables/pt-BR/source-snapshot.json. Essa fotografia não afirma que existe um commit Git; a sessão não pôde criá-lo.')
(OUT/'guia-tecnico-completo.md').write_text(source)
d=ns['base']('HEDAMI Pricing Advisor','Guia pessoal de estudo: solução, arquitetura, agentes e código')
ns['render'](d,source,True)
replacements={
'Prepared for Heron Dantas':'Preparado para Heron Dantas',
'Detailed study document and English portfolio reference\nEvidence snapshot: 24 October 2026\nVersion 1.0 — revised after documented author self-review':
'Material pessoal de estudo — não é a apresentação final\nEvidências de referência: 24 de outubro de 2026\nEdição em português do Brasil — versão 1.0',
'Purpose: understand and defend the implementation.\nStatus: executed console prototype; fictional business inputs; recommendation only.\nNo claim of production readiness or demonstrated superiority of extra agents.':
'Objetivo: entender o que foi construído e preparar perguntas para discussão.\nUso pessoal de Heron Dantas; ainda não é material final de apresentação.\nProtótipo executado no console, com dados fictícios e recomendações sem alteração de preço.',
'The code, evidence and limitations are explained together. Runtime procedures and development skills are distinguished explicitly. This document does not contain an API key.':
'Código, evidências e limitações são explicados em conjunto. Skills de desenvolvimento e procedimentos da aplicação são diferenciados. Nenhuma chave da API foi incluída.',
'Chapter navigation':'Navegação pelos capítulos',
'Click a title below or use Word’s Navigation Pane. Page numbers are generated by Word during layout.':
'Clique em um título ou use o Painel de Navegação do Word. A numeração de páginas é calculada ao abrir o documento.'
}
def paragraphs(document):
    yield from document.paragraphs
    for t in document.tables:
        for row in t.rows:
            for cell in row.cells:yield from cell.paragraphs
for p in paragraphs(d):
    if p.text in replacements:p.text=replacements[p.text]
    if p.text.startswith('Fonte do trecho:'):
        p.style='Caption';p.paragraph_format.keep_with_next=True
for section in d.sections:
    for p in section.footer.paragraphs:
        for run in p.runs:
            if 'Technical study guide' in run.text:
                run.text=run.text.replace('Technical study guide  •  Evidence snapshot: 24 Oct 2026','Guia pessoal de estudo  •  Evidências: 24 out. 2026')
lang=d.styles['Normal'].element.get_or_add_rPr().find(qn('w:lang'))
lang.set(qn('w:val'),'pt-BR')
d.core_properties.language='pt-BR'
d.core_properties.title='HEDAMI Pricing Advisor — guia pessoal em português do Brasil'
d.core_properties.subject='Estudo pessoal da solução, arquitetura, código, agentes, harness, skills e avaliações'
d.core_properties.author='Heron Dantas — elaborado com assistência do Codex'
d.core_properties.comments='Uso pessoal para entendimento. Não é apresentação final. Tradução dos procedimentos é explicativa; runtime original preservado.'
path=OUT/'HEDAMI-Pricing-Advisor-Guia-Pessoal-pt-BR.docx';d.save(path)
assert hashlib.sha256(english.read_bytes()).hexdigest()==english_hash
with zipfile.ZipFile(path) as z:assert z.testzip() is None
print(json.dumps({'arquivo':str(path),'palavrasDaFonte':len(source.split()),'capitulos':len(re.findall(r'^# ',source,re.M)),'diagramas':len(d.inline_shapes),'tabelas':len(d.tables),'trechosReais':len(ns['SNIPS']),'versaoInglesPreservada':True},ensure_ascii=False,indent=2))
